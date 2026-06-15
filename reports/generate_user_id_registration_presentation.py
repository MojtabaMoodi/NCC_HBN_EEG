#!/usr/bin/env python3
"""Generate PowerPoint: LaBraM user registration results only (1s, 2s, 4s).

Data source: registering_users_results/pipeline_{1s,2s,4s}_active_passive/
"""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

EEG_ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = EEG_ROOT / "reports" / "presentations"
OUT_DIR.mkdir(parents=True, exist_ok=True)

REG_ROOT = EEG_ROOT / "registering_users_results"
REG_SOURCE = "registering_users_results"

SEGMENTS = ("1s", "2s", "4s")
TASKS = ("active", "passive")
EXPECTED_TABLE_ROWS = len(SEGMENTS) * len(TASKS)

# Light slide background, high-contrast table
BG = RGBColor(0xFF, 0xFF, 0xFF)
TITLE_COLOR = RGBColor(0x0F, 0x17, 0x2A)
BODY_COLOR = RGBColor(0x1E, 0x29, 0x3B)
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
HEADER_FILL = RGBColor(0x1E, 0x40, 0xAF)
HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
ROW_FILL_A = RGBColor(0xFF, 0xFF, 0xFF)
ROW_FILL_B = RGBColor(0xE0, 0xE7, 0xFF)

REGISTRATION_METHOD_NOTE = (
    "User registration tests whether new EEG users can be enrolled from their segments and "
    "later recognized, while strangers are rejected. Only held-out unknown participants are used. "
    "Per fold: 50% of unknown users are enrolled; 50% remain unregistered strangers. "
    "For each enrolled user, 70% of segments build an L2-normalized prototype (mean ArcFace "
    "embedding from the LaBraM model). Each probe is matched to the nearest prototype; "
    "segments with match score 1/(1 + d) below threshold τ are labeled UNKNOWN. "
    "Threshold τ is swept (501 points) to maximize weighted_recall_balance on registered probes "
    "plus all stranger segments. Open-world accuracy is the fraction of correct decisions "
    "(correct enrolled user or UNKNOWN). Values are 5-fold mean ± std from registering_users_results."
)


def load_json(path: Path) -> dict:
    with path.open() as f:
        return json.load(f)


def load_registration_manifest(segment: str) -> dict:
    path = REG_ROOT / f"pipeline_{segment}_active_passive" / "registration_pipeline_manifest.json"
    if not path.is_file():
        raise FileNotFoundError(f"Missing registration manifest: {path}")
    return load_json(path)


def pct(mean: float, std: float | None = None) -> str:
    if std is None:
        return f"{mean * 100:.1f}%"
    return f"{mean * 100:.1f}% ± {std * 100:.1f}%"


def set_slide_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def set_cell_text(
    cell,
    text: str,
    *,
    header: bool = False,
    row_index: int = 0,
) -> None:
    """Write exactly one paragraph per cell (avoids blank extra rows in PowerPoint)."""
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = False
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf.paragraphs[0]
    p.text = str(text)
    p.alignment = PP_ALIGN.CENTER
    p.font.bold = header
    p.font.size = Pt(13 if header else 12)
    p.font.color.rgb = HEADER_TEXT if header else BODY_COLOR

    if header:
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_FILL
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = ROW_FILL_B if row_index % 2 else ROW_FILL_A


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(12.0), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TITLE_COLOR

    box2 = slide.shapes.add_textbox(Inches(0.6), Inches(3.5), Inches(12.0), Inches(1.0))
    p2 = box2.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(20)
    p2.font.color.rgb = ACCENT


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.0), Inches(0.8))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(28)
    tp.font.bold = True
    tp.font.color.rgb = TITLE_COLOR

    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)
        p.font.color.rgb = BODY_COLOR
        p.space_after = Pt(8)


def add_registration_table_slide(
    prs: Presentation,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    note: str | None = None,
    top: float = 1.0,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)

    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.6))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(22)
    tp.font.bold = True
    tp.font.color.rgb = TITLE_COLOR

    nrows = len(rows) + 1
    ncols = len(headers)
    row_h_in = 0.58
    table_h = Inches(row_h_in * nrows)
    table_top = Inches(top)

    shape = slide.shapes.add_table(nrows, ncols, Inches(0.4), table_top, Inches(12.5), table_h)
    table = shape.table

    # Do not set per-column widths (causes layout glitches / empty rows in PowerPoint).
    for r in range(nrows):
        table.rows[r].height = Inches(row_h_in)

    for j, h in enumerate(headers):
        set_cell_text(table.cell(0, j), h, header=True)

    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            set_cell_text(table.cell(i, j), val, row_index=i)

    if note:
        note_y = top + row_h_in * nrows + 0.12
        nbox = slide.shapes.add_textbox(Inches(0.5), Inches(note_y), Inches(12.3), Inches(0.45))
        np_ = nbox.text_frame.paragraphs[0]
        np_.text = note
        np_.font.size = Pt(10)
        np_.font.color.rgb = BODY_COLOR
        np_.font.italic = True


def add_note_slide(prs: Presentation, title: str, body: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.0), Inches(0.8))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = TITLE_COLOR

    bbox = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(11.8), Inches(5.8))
    tf = bbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body
    p.font.size = Pt(16)
    p.font.color.rgb = BODY_COLOR
    p.line_spacing = 1.25


def add_image_slide(prs: Presentation, title: str, image_path: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.7))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(24)
    tp.font.bold = True
    tp.font.color.rgb = TITLE_COLOR
    slide.shapes.add_picture(str(image_path), Inches(0.8), Inches(1.0), width=Inches(11.5))


def make_bar_chart(
    out_path: Path,
    title: str,
    labels: list[str],
    series: dict[str, list[float]],
) -> None:
    x = np.arange(len(labels))
    n_series = len(series)
    width = 0.8 / max(n_series, 1)
    fig, ax = plt.subplots(figsize=(10, 5), facecolor="#ffffff")
    ax.set_facecolor("#ffffff")
    colors = ["#1d4ed8", "#047857"]
    for i, (name, vals) in enumerate(series.items()):
        offset = (i - (n_series - 1) / 2) * width
        bars = ax.bar(x + offset, vals, width, label=name, color=colors[i % len(colors)], edgecolor="white")
        for bar in bars:
            h = bar.get_height()
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                h + 0.4,
                f"{h:.1f}%",
                ha="center",
                va="bottom",
                color="#0f172a",
                fontsize=11,
                fontweight="bold",
            )
    ax.set_xticks(x)
    ax.set_xticklabels(labels, color="#0f172a", fontsize=12)
    ax.set_ylabel("Open-world accuracy (%)", color="#0f172a", fontsize=12)
    ax.set_title(title, color="#0f172a", fontsize=15, fontweight="bold", pad=14)
    ax.tick_params(colors="#334155", labelsize=11)
    ax.set_ylim(70, 90)
    ax.legend(facecolor="#f8fafc", edgecolor="#cbd5e1", labelcolor="#0f172a", fontsize=11)
    ax.grid(axis="y", linestyle="--", alpha=0.45, color="#94a3b8")
    for spine in ax.spines.values():
        spine.set_color("#cbd5e1")
    fig.tight_layout()
    fig.savefig(out_path, dpi=150, facecolor="#ffffff")
    plt.close(fig)


def build_registration_rows() -> list[dict]:
    """One dict per (window, task) from registering_users_results only."""
    records: list[dict] = []
    for seg in SEGMENTS:
        manifest = load_registration_manifest(seg)
        for task in TASKS:
            agg = manifest[task]["best_config"]["aggregate"]
            records.append(
                {
                    "window": seg,
                    "task": task,
                    "open_world": pct(
                        agg["open_world_accuracy_mean"], agg["open_world_accuracy_std"]
                    ),
                    "balanced": pct(
                        agg["open_world_balanced_accuracy_mean"],
                        agg["open_world_balanced_accuracy_std"],
                    ),
                    "unknown_recall": pct(agg["unknown_recall_mean"], agg["unknown_recall_std"]),
                    "registered_recall": pct(
                        agg["mean_registered_recall_mean"], agg["mean_registered_recall_std"]
                    ),
                    "probe_acc": pct(
                        agg["registered_probe_accuracy_mean"],
                        agg["registered_probe_accuracy_std"],
                    ),
                    "open_world_val": agg["open_world_accuracy_mean"] * 100,
                }
            )
    return records


def records_to_table(records: list[dict], columns: list[str]) -> list[list[str]]:
    key_map = {
        "Window": "window",
        "Task": "task",
        "Open-world acc.": "open_world",
        "Balanced acc.": "balanced",
        "Unknown recall": "unknown_recall",
        "Registered recall": "registered_recall",
        "Probe acc.": "probe_acc",
    }
    return [[str(r[key_map[c]]) for c in columns] for r in records]


def build_summary_bullets(records: list[dict]) -> list[str]:
    best = max(records, key=lambda r: r["open_world_val"])
    return [
        f"Best open-world: {best['window']} {best['task']} — {best['open_world']} "
        f"(unknown recall {best['unknown_recall']}).",
        "4s active: highest open-world (~83.5%); passive is ~5 points lower at each window.",
        "1s active ~81%; 2s active ~82%; registration is strong at all segment lengths.",
        "Active EEG outperforms passive for registration (1s, 2s, 4s).",
        f"All metrics from {REG_SOURCE} (registration pipeline, 5-fold CV).",
    ]


def collect_slide_text(prs: Presentation) -> str:
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def main() -> None:
    charts_dir = OUT_DIR / "_charts"
    charts_dir.mkdir(exist_ok=True)

    records = build_registration_rows()
    if len(records) != EXPECTED_TABLE_ROWS:
        raise RuntimeError(f"Expected {EXPECTED_TABLE_ROWS} rows, got {len(records)}")

    reg_active = {r["window"]: r["open_world_val"] for r in records if r["task"] == "active"}
    reg_passive = {r["window"]: r["open_world_val"] for r in records if r["task"] == "passive"}

    chart_path = charts_dir / "registration_open_world.png"
    make_bar_chart(
        chart_path,
        "User registration — open-world accuracy (5-fold CV)",
        list(SEGMENTS),
        {
            "Active": [reg_active[s] for s in SEGMENTS],
            "Passive": [reg_passive[s] for s in SEGMENTS],
        },
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "User Registration Results",
        "LaBraM · EEG window sizes 1s, 2s, 4s · registering_users_results",
    )

    add_bullet_slide(
        prs,
        "User registration overview",
        [
            "Goal: enroll new EEG users and recognize their probes; reject unregistered strangers.",
            f"Data source: {REG_SOURCE}/pipeline_{{1s,2s,4s}}_active_passive/.",
            "LaBraM + ArcFace embeddings; prototypes built from enrollment segments.",
            "5-fold cross-validation; values are mean ± std.",
            "Window sizes 1s, 2s, 4s — active and passive tasks reported separately.",
        ],
    )

    add_bullet_slide(
        prs,
        "How registration metrics are calculated",
        [
            "1. Evaluate only held-out unknown participants.",
            "2. Enroll 50% of unknown users; 50% act as unregistered strangers.",
            "3. Per enrolled user: 70% of segments → prototype (mean L2-normalized embedding).",
            "4. Probes: nearest prototype; reject as UNKNOWN if match score 1/(1+d) < τ.",
            "5. Threshold τ tuned via grid search (weighted_recall_balance).",
            "6. Open-world accuracy = correct enrolled-user match + correct UNKNOWN on strangers.",
        ],
    )

    add_note_slide(prs, "Registration methodology", REGISTRATION_METHOD_NOTE)

    # Table 1: primary accuracy metrics (5 columns — easier to read)
    add_registration_table_slide(
        prs,
        "Registration results — accuracy (1s, 2s, 4s)",
        ["Window", "Task", "Open-world acc.", "Balanced acc.", "Unknown recall"],
        records_to_table(
            records,
            ["Window", "Task", "Open-world acc.", "Balanced acc.", "Unknown recall"],
        ),
        note=f"Source: {REG_SOURCE}/pipeline_{{1s,2s,4s}}_active_passive/registration_pipeline_manifest.json",
        top=1.0,
    )

    # Table 2: recall / probe metrics (same 6 rows, no empty rows)
    add_registration_table_slide(
        prs,
        "Registration results — per-user recall (1s, 2s, 4s)",
        ["Window", "Task", "Registered recall", "Registered probe acc."],
        records_to_table(records, ["Window", "Task", "Registered recall", "Probe acc."]),
        top=1.0,
    )

    add_image_slide(
        prs,
        "Registration open-world accuracy — 1s, 2s, 4s",
        chart_path,
    )

    add_bullet_slide(prs, "Summary", build_summary_bullets(records))

    out_pptx = OUT_DIR / "user_registration_results.pptx"
    prs.save(str(out_pptx))

    # Verify tables and registration-only wording
    check = Presentation(str(out_pptx))
    all_text = collect_slide_text(check).lower()
    if "identification" in all_text:
        raise RuntimeError("Presentation must not mention user identification")

    for si, slide in enumerate(check.slides, 1):
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            t = shape.table
            for r in range(1, len(t.rows)):
                for c in range(len(t.columns)):
                    txt = t.cell(r, c).text.strip()
                    if not txt:
                        raise RuntimeError(f"Empty cell on slide {si} row {r} col {c}")

    # Remove legacy combined deck if present
    legacy = OUT_DIR / "user_identification_and_registration_results.pptx"
    if legacy.is_file():
        legacy.unlink()

    print(f"Saved: {out_pptx} ({len(records)} data rows per table)")


if __name__ == "__main__":
    main()
